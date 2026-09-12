"""Kiểm thử bộ sinh khung báo cáo Word.

VÌ SAO SINH BẰNG MÃ THAY VÌ GÕ TAY
===================================
Đề bài kèm 7 ảnh hướng dẫn định dạng Word rất cụ thể (Times New Roman 14, lề
2–2–3–2, viền trang bìa). Gõ tay thì mỗi lần chỉ số thay đổi lại phải sửa lại
báo cáo, và sai lề thì không ai phát hiện cho tới lúc nộp.

Sinh bằng mã thì: định dạng được KIỂM THỬ, còn số liệu LẤY THẲNG từ
``eval/ket_qua_danh_gia.json`` nên không bao giờ lệch với thực đo.
"""
from pathlib import Path

import pytest

docx = pytest.importorskip("docx", reason="Cần gói tuỳ chọn 'bao-cao'")

from docx.shared import Cm, Pt  # noqa: E402

from tra_cuu_gtdb.bao_cao.tao_bao_cao import DINH_DANG, dung_bao_cao  # noqa: E402

GOC = Path(__file__).resolve().parent.parent


@pytest.fixture(scope="module")
def tai_lieu():
    return dung_bao_cao(GOC)


class TestDinhDangTheoHuongDan:
    """Bốn thiết lập rút ra từ 7 ảnh hướng dẫn — sai một cái là trừ điểm trình bày."""

    #: Word lưu lề theo twip (1 twip = 635 EMU), nên Cm(2) = 720000 EMU đọc lại
    #: thành 720090. Sai số 0,00025 cm là ĐỘ MỊN CỦA ĐỊNH DẠNG, không phải lỗi —
    #: so sánh bằng nhau tuyệt đối ở đây là đòi hỏi bất khả thi.
    SAI_SO_EMU = 635

    def test_le_trang_dung_2_2_3_2_cm(self, tai_lieu):
        khu = tai_lieu.sections[0]
        for ten, thuc, mong_doi in [
            ("trên", khu.top_margin, Cm(2)),
            ("dưới", khu.bottom_margin, Cm(2)),
            ("trái", khu.left_margin, Cm(3)),
            ("phải", khu.right_margin, Cm(2)),
        ]:
            assert abs(thuc - mong_doi) <= self.SAI_SO_EMU, (
                f"Lề {ten}: {thuc / 360000:.3f} cm, cần {mong_doi / 360000:.0f} cm")

    def test_font_mac_dinh_times_new_roman_co_14(self, tai_lieu):
        kieu = tai_lieu.styles["Normal"].font
        assert kieu.name == "Times New Roman"
        assert kieu.size == Pt(14)

    def test_huong_trang_dung(self, tai_lieu):
        from docx.enum.section import WD_ORIENT
        assert tai_lieu.sections[0].orientation == WD_ORIENT.PORTRAIT

    def test_hang_so_dinh_dang_khop_huong_dan(self):
        assert DINH_DANG["font"] == "Times New Roman"
        assert DINH_DANG["co_chu"] == 14
        assert DINH_DANG["le_cm"] == {"tren": 2, "duoi": 2, "trai": 3, "phai": 2}


class TestNoiDung:
    def test_co_du_cac_muc_bat_buoc(self, tai_lieu):
        van_ban = "\n".join(p.text for p in tai_lieu.paragraphs)
        for muc in ["Giới thiệu", "Cơ sở tri thức", "Thiết kế giải pháp",
                    "Thực nghiệm", "Hạn chế", "Kết luận"]:
            assert muc in van_ban, f"Thiếu mục {muc!r}"

    def test_trang_bia_neu_du_thong_tin_dinh_danh(self, tai_lieu):
        van_ban = "\n".join(p.text for p in tai_lieu.paragraphs)
        for x in ["CS106", "Nhóm 7", "Nguyễn Đình Hiển", "Đề tài 4"]:
            assert x in van_ban, f"Trang bìa thiếu {x!r}"

    def test_liet_ke_du_bay_thanh_vien(self, tai_lieu):
        import json
        with (GOC / "docs" / "thanh_vien.json").open(encoding="utf-8") as f:
            tv = json.load(f)["thanh_vien"]
        van_ban = "\n".join(p.text for p in tai_lieu.paragraphs)
        van_ban += "\n".join(o.text for b in tai_lieu.tables
                             for h in b.rows for o in h.cells)
        assert len(tv) == 7
        for nguoi in tv:
            assert nguoi["mssv"] in van_ban, f"Thiếu MSSV {nguoi['mssv']}"


class TestSoLieuLayTuKetQuaThuc:
    """Số trong báo cáo phải BẰNG số đo được, không được gõ cứng."""

    def test_top1_khop_ket_qua_danh_gia(self, tai_lieu):
        import json
        with (GOC / "eval" / "ket_qua_danh_gia.json").open(encoding="utf-8") as f:
            th = json.load(f)["tong_hop"]
        mong_doi = f"{th['top1']:.2%}".replace(".", ",")
        moi_o = "\n".join(o.text for b in tai_lieu.tables
                          for h in b.rows for o in h.cells)
        assert mong_doi in moi_o, f"Không thấy Top-1 {mong_doi} trong bảng chỉ số"

    def test_khong_go_cung_chi_so_trong_ma_nguon(self):
        """Chốt chặn: số liệu phải đọc từ JSON, không nằm chết trong mã."""
        nguon = (GOC / "src" / "tra_cuu_gtdb" / "bao_cao" / "tao_bao_cao.py"
                 ).read_text(encoding="utf-8")
        for so in ["76,67", "0,8384", "95,83"]:
            assert so not in nguon, f"Chỉ số {so} bị gõ cứng trong mã nguồn"
