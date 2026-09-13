"""Kiểm thử bộ sinh khung báo cáo Word.

VÌ SAO SINH BẰNG MÃ THAY VÌ GÕ TAY
===================================
Đề bài kèm 7 ảnh hướng dẫn định dạng Word rất cụ thể (Times New Roman 14, lề
2–2–3–2, viền trang bìa). Gõ tay thì mỗi lần chỉ số thay đổi lại phải sửa lại
báo cáo, và sai lề thì không ai phát hiện cho tới lúc nộp.

Sinh bằng mã thì: định dạng được KIỂM THỬ, còn số liệu LẤY THẲNG từ
``eval/ket_qua_danh_gia.json`` nên không bao giờ lệch với thực đo.
"""
from pathlib import Path

import pytest

docx = pytest.importorskip("docx", reason="Cần gói tuỳ chọn 'bao-cao'")

from docx.shared import Cm, Pt  # noqa: E402

from traffic_law.report.builder import FORMAT, build_report  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture(scope="module")
def document():
    return build_report(ROOT)


class TestFormatFollowsGuide:
    """Bốn thiết lập rút ra từ 7 ảnh hướng dẫn — sai một cái là trừ điểm trình bày."""

    #: Word lưu lề theo twip (1 twip = 635 EMU), nên Cm(2) = 720000 EMU đọc lại
    #: thành 720090. Sai số 0,00025 cm là ĐỘ MỊN CỦA ĐỊNH DẠNG, không phải lỗi —
    #: so sánh bằng nhau tuyệt đối ở đây là đòi hỏi bất khả thi.
    SAI_SO_EMU = 635

    def test_margins_are_2_2_3_2_cm(self, document):
        khu = document.sections[0]
        for name, thuc, expected in [
            ("trên", khu.top_margin, Cm(2)),
            ("dưới", khu.bottom_margin, Cm(2)),
            ("trái", khu.left_margin, Cm(3)),
            ("phải", khu.right_margin, Cm(2)),
        ]:
            assert abs(thuc - expected) <= self.SAI_SO_EMU, (
                f"Lề {name}: {thuc / 360000:.3f} cm, cần {expected / 360000:.0f} cm")

    def test_default_font_is_times_new_roman_14(self, document):
        kind = document.styles["Normal"].font
        assert kind.name == "Times New Roman"
        assert kind.size == Pt(14)

    def test_page_orientation_is_portrait(self, document):
        from docx.enum.section import WD_ORIENT
        assert document.sections[0].orientation == WD_ORIENT.PORTRAIT

    def test_format_constants_match_the_guide(self):
        assert FORMAT["font"] == "Times New Roman"
        assert FORMAT["co_chu"] == 14
        assert FORMAT["le_cm"] == {"tren": 2, "duoi": 2, "trai": 3, "phai": 2}


class TestContent:
    def test_has_all_required_sections(self, document):
        documents = "\n".join(p.text for p in document.paragraphs)
        for item in ["Giới thiệu", "Cơ sở tri thức", "Thiết kế giải pháp",
                    "Thực nghiệm", "Hạn chế", "Kết luận"]:
            assert item in documents, f"Thiếu mục {item!r}"

    def test_cover_page_states_all_identifying_info(self, document):
        documents = "\n".join(p.text for p in document.paragraphs)
        for x in ["CS106", "Nhóm 7", "Nguyễn Đình Hiển", "Đề tài 4"]:
            assert x in documents, f"Trang bìa thiếu {x!r}"

    def test_lists_all_seven_members(self, document):
        import json
        with (ROOT / "docs" / "thanh_vien.json").open(encoding="utf-8") as f:
            tv = json.load(f)["thanh_vien"]
        documents = "\n".join(p.text for p in document.paragraphs)
        documents += "\n".join(o.text for b in document.tables
                             for h in b.rows for o in h.cells)
        assert len(tv) == 7
        for nguoi in tv:
            assert nguoi["mssv"] in documents, f"Thiếu MSSV {nguoi['mssv']}"


class TestFiguresFromRealResults:
    """Số trong báo cáo phải BẰNG số đo được, không được gõ cứng."""

    def test_top1_matches_evaluation_result(self, document):
        import json
        with (ROOT / "eval" / "ket_qua_danh_gia.json").open(encoding="utf-8") as f:
            th = json.load(f)["summary"]
        expected = f"{th['top1']:.2%}".replace(".", ",")
        moi_o = "\n".join(o.text for b in document.tables
                          for h in b.rows for o in h.cells)
        assert expected in moi_o, f"Không thấy Top-1 {expected} trong bảng chỉ số"

    def test_metrics_not_hardcoded_in_source(self):
        """Chốt chặn: số liệu phải đọc từ JSON, không nằm chết trong mã."""
        source = (ROOT / "src" / "traffic_law" / "report" / "builder.py"
                 ).read_text(encoding="utf-8")
        for so in ["76,67", "0,8384", "95,83"]:
            assert so not in source, f"Chỉ số {so} bị gõ cứng trong mã nguồn"
